"""Read-only inventory helper for the Mandros Blind reference files.

This script intentionally never modifies the supplied workbooks or documents.
It prints compact JSON summaries suitable for building and testing the local
validation contract.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber
from docx import Document


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def safe_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def inspect_workbook(
    path: Path,
    max_rows: int = 120,
    max_cols: int = 40,
    sheet_names: set[str] | None = None,
    summary_only: bool = False,
) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        if sheet_names and sheet.title not in sheet_names:
            continue
        nonempty: list[dict[str, Any]] = []
        for row in sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, max_rows),
            min_col=1,
            max_col=min(sheet.max_column, max_cols),
        ):
            values = [safe_value(cell.value) for cell in row]
            while values and values[-1] in (None, ""):
                values.pop()
            if any(value not in (None, "") for value in values):
                nonempty.append({"row": row[0].row, "values": values})
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "nonempty_rows": [] if summary_only else nonempty,
            }
        )
    return {
        "path": str(path),
        "sha256": sha256(path),
        "sheet_count": len(workbook.worksheets),
        "sheets": sheets,
    }


def inspect_docx(path: Path) -> dict[str, Any]:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            rows.append(
                {
                    "row": row_index,
                    "values": [cell.text for cell in row.cells],
                }
            )
        tables.append({"table": table_index, "rows": rows})
    return {
        "path": str(path),
        "sha256": sha256(path),
        "paragraphs": paragraphs,
        "tables": tables,
    }


def inspect_pdf(path: Path) -> dict[str, Any]:
    pages = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            pages.append({"page": page_number, "text": text})
    return {
        "path": str(path),
        "sha256": sha256(path),
        "page_count": len(pages),
        "pages": pages,
    }


def inspect_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    return {
        "path": str(path),
        "sha256": sha256(path),
        "payload": payload,
    }


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("kind", choices=("xlsx", "docx", "pdf", "json"))
    parser.add_argument("path", type=Path)
    parser.add_argument("--max-rows", type=int, default=120)
    parser.add_argument("--max-cols", type=int, default=40)
    parser.add_argument("--sheet", action="append", dest="sheets")
    parser.add_argument("--summary-only", action="store_true")
    args = parser.parse_args()

    if args.kind == "xlsx":
        result = inspect_workbook(
            args.path,
            args.max_rows,
            args.max_cols,
            set(args.sheets or []),
            args.summary_only,
        )
    elif args.kind == "docx":
        result = inspect_docx(args.path)
    elif args.kind == "pdf":
        result = inspect_pdf(args.path)
    else:
        result = inspect_json(args.path)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
