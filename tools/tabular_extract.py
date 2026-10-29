"""Extract selected columns from a flat worksheet for compact inspection."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("sheet")
    parser.add_argument("headers", nargs="+")
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.path, data_only=False, read_only=True)
    sheet = workbook[args.sheet]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    positions = {str(value): index for index, value in enumerate(headers) if value is not None}
    missing = [header for header in args.headers if header not in positions]
    if missing:
        raise SystemExit(f"Missing headers: {missing}; available: {list(positions)}")
    output = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        record = {header: values[positions[header]] for header in args.headers}
        if any(value not in (None, "") for value in record.values()):
            output.append({"excel_row": excel_row, **record})
    print(json.dumps(output, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
