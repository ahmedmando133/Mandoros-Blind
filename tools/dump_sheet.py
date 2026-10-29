"""Print selected spreadsheet cells with stable Excel coordinates."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("sheet")
    parser.add_argument("--rows", default="1:999999")
    parser.add_argument("--cols", default="1:999999")
    parser.add_argument("--skip-empty", action="store_true")
    args = parser.parse_args()

    row_start, row_end = (int(part) for part in args.rows.split(":", 1))
    col_start, col_end = (int(part) for part in args.cols.split(":", 1))
    workbook = openpyxl.load_workbook(args.path, data_only=False, read_only=False)
    sheet = workbook[args.sheet]
    output = []
    for row in sheet.iter_rows(
        min_row=row_start,
        max_row=min(row_end, sheet.max_row),
        min_col=col_start,
        max_col=min(col_end, sheet.max_column),
    ):
        cells = [
            {"cell": cell.coordinate, "value": cell.value}
            for cell in row
            if cell.value not in (None, "")
        ]
        if cells or not args.skip_empty:
            output.append({"row": row[0].row, "cells": cells})
    print(json.dumps(output, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
