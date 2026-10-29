from __future__ import annotations

import json
import subprocess
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

REFERENCE_14_SHEET_WORKBOOKS = {
    "orlistat": ROOT / "outputs" / "studies" / "orlistat-rct-2024-34aa622cbc7d" / "mandros_input_corrected.xlsx",
    "dapagliflozin": ROOT / "outputs" / "studies" / "private-dapagliflozin-trial-daba-dm_extracted-bb3afe6702e6" / "daba-dm_extracted_corrected.xlsx",
    "semaglutide": ROOT / "outputs" / "studies" / "private-semaglutide-obesity-trial-semaglutide-obesity-without-dm_extracted-1676c841b082" / "semaglutide-obesity-without-dm_extracted_corrected.xlsx",
    "simvastatin": ROOT / "outputs" / "studies" / "private-simvastatin-simvastatin_extracted-81abce43466b" / "simvastatin_extracted_corrected.xlsx",
    "empagliflozin": ROOT / "outputs" / "studies" / "private-empagliflozin-2-empa-vs-placebo_extracted-45ee13aae0e3" / "empa-vs-placebo_extracted_corrected.xlsx",
    "empa_metsu": ROOT / "outputs" / "studies" / "private-empa-plus-met-empa-vs-placebo-add-on-metformin-rct3-_extracted-ad7cbe59c74a" / "empa-vs-placebo-add-on-metformin-rct3-_extracted_corrected.xlsx",
    "empa_rct4": ROOT / "outputs" / "studies" / "private-empagliflozin-empa_vs_placebo_2doses__rct4_extracted-7-4ac9c0b7b1ae" / "empa_vs_placebo_2doses__rct4_extracted-7_corrected.xlsx",
}


class WebExportParityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        node_script = (
            "global.window = {};"
            "global.document = { addEventListener: () => {}, getElementById: () => null, querySelectorAll: () => [] };"
            "require('./benchmarks.js');"
            "require('./app.js');"
            "const out = {};"
            "for (const [key, study] of Object.entries(window.MB_BENCHMARKS)) {"
            "  const m = window.MB_ENGINE.build14SheetMatrices(study.bundle);"
            "  out[key] = {"
            "    export_filename: study.export_filename,"
            "    sheets: Object.keys(m),"
            "    headers: Object.fromEntries(Object.entries(m).map(([s, v]) => [s, v.aoa[0]])),"
            "    row_counts: Object.fromEntries(Object.entries(m).map(([s, v]) => [s, v.aoa.length]))"
            "  };"
            "}"
            "console.log(JSON.stringify(out));"
        )
        raw = subprocess.check_output(["node", "-e", node_script], cwd=ROOT, text=True, encoding="utf-8")
        cls.generated = json.loads(raw)

    def test_all_eight_studies_have_exact_14_sheets_and_column_parity(self) -> None:
        expected_14_sheets = [
            "Master Extraction",
            "Studies",
            "Arms",
            "Eligibility",
            "Baseline",
            "Outcomes",
            "Statistical Tests",
            "Document Inventory",
            "Source Conflicts",
            "Phase 2 Readiness",
            "Statistical Methods",
            "Provenance",
            "Validation Log",
            "Quality Assessment",
        ]
        self.assertIn("metformin", self.generated)
        self.assertEqual(expected_14_sheets, self.generated["metformin"]["sheets"])
        self.assertEqual("Mandroos_Blind_Metformin_DPPOS.xlsx", self.generated["metformin"]["export_filename"])

        for key, ref_path in REFERENCE_14_SHEET_WORKBOOKS.items():
            self.assertIn(key, self.generated, f"Missing study {key} in benchmarks.js")
            gen = self.generated[key]
            self.assertEqual(expected_14_sheets, gen["sheets"])
            self.assertTrue(gen["export_filename"].startswith("Mandroos_Blind_"))
            self.assertTrue(gen["export_filename"].endswith(".xlsx"))

            wb = load_workbook(ref_path, data_only=False, read_only=True)
            try:
                self.assertEqual(wb.sheetnames, gen["sheets"], f"Sheet order mismatch for {key}")
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows = list(ws.iter_rows(values_only=True))
                    ref_header = list(rows[0]) if rows else []
                    while ref_header and ref_header[-1] is None:
                        ref_header.pop()
                    gen_header = list(gen["headers"][sheet_name])
                    while gen_header and gen_header[-1] is None:
                        gen_header.pop()
                    if sheet_name == "Studies" and key == "orlistat":
                        # standards/orlistat_workbook_dataset.json includes source-checked study keys
                        self.assertIn("study_id", gen_header)
                        continue
                    self.assertEqual(ref_header, gen_header, f"Header mismatch in {key} -> {sheet_name}")
                    if key != "orlistat":
                        self.assertEqual(len(rows), gen["row_counts"][sheet_name], f"Row count mismatch in {key} -> {sheet_name}")
            finally:
                wb.close()

    def test_index_html_has_three_mode_dual_upload_banner_and_mandoros_link(self) -> None:
        html = (ROOT / "index.html").read_text(encoding="utf-8")
        self.assertIn("https://mandoros.onrender.com/", html)
        self.assertIn("Recommended for 100% Accuracy &amp; Fastest Performance: Upload BOTH the Original Trial PDF and the Extracted Excel from", html)
        self.assertIn("Upload Original Trial PDF", html)
        self.assertIn("Upload Extracted Excel from MANDOROS", html)


if __name__ == "__main__":
    unittest.main()
