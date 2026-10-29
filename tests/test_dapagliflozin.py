import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.validator import validate_workbook


ROOT = Path(__file__).resolve().parents[1]
DATASET = ROOT / "standards" / "dapagliflozin_workbook_dataset.json"
PDF = ROOT / "data" / "private" / "dapagliflozin trial" / "dabagliflozin RCT1.pdf"
MANDROS = ROOT / "data" / "private" / "dapagliflozin trial" / "daba DM_extracted.xlsx"
OUTPUT = ROOT / "outputs" / "studies" / "private-dapagliflozin-trial-daba-dm_extracted-bb3afe6702e6" / "daba-dm_extracted_corrected.xlsx"


class DapagliflozinTruthBundleTests(unittest.TestCase):
    def test_dataset_has_complete_main_article_tables(self):
        data = json.loads(DATASET.read_text(encoding="utf-8"))
        self.assertEqual(data["dataset_id"], "dapagliflozin-monotherapy-dc10-0612-v1")
        self.assertEqual(len(data["arms"]), 9)
        self.assertEqual(len(data["baseline"]), 72)
        self.assertEqual(len(data["outcomes"]), 257)
        self.assertEqual(len(data["statistical_tests"]), 6)
        arm_ids = {row["arm_id"] for row in data["arms"]}
        self.assertFalse(any(row.get("arm_id") and row["arm_id"] not in arm_ids for row in data["baseline"] + data["outcomes"]))
        self.assertTrue(all("p_value" not in row for row in data["outcomes"]))
        self.assertTrue(any(row["outcome"] == "events suggestive of genital infections" for row in data["outcomes"]))
        sex_rows = [row for row in data["baseline"] if row["characteristic"] == "sex"]
        self.assertEqual(len(sex_rows), 18)
        self.assertTrue(all(row["measure"] == "sample_proportion" for row in sex_rows))
        self.assertTrue(all(row["event_count"] is not None and row["denominator"] is not None for row in sex_rows))
        safety_keys = [(row["outcome"], row.get("arm_id"), row.get("timepoint"), row.get("subgroup_value")) for row in data["outcomes"] if row["domain"] == "safety" and row.get("arm_id") and row["outcome"] != "nocturia"]
        self.assertEqual(len(safety_keys), len(set(safety_keys)))

    def test_inclusive_exclusions_and_morning_a1c_subgroup(self):
        data = json.loads(DATASET.read_text(encoding="utf-8"))
        criteria = {row["criterion_id"]: row for row in data["eligibility"]}
        for criterion_id in ("ELIG_EXC_CREAT_MEN", "ELIG_EXC_CREAT_WOMEN", "ELIG_EXC_CK", "ELIG_EXC_SBP", "ELIG_EXC_DBP"):
            self.assertTrue(criteria[criterion_id]["operator_eq"])
            self.assertTrue(criteria[criterion_id]["operator_gt"])
        subgroup_rows = [row for row in data["outcomes"] if row["id"].startswith("EFF_A1C_GE9_")]
        self.assertEqual({row["arm_id"] for row in subgroup_rows}, {"ARM_PBO_AM", "ARM_DAPA_2_5_AM", "ARM_DAPA_5_AM", "ARM_DAPA_10_AM"})
        self.assertTrue(all(row["analysis_population"] == "primary_morning_dose_cohort_with_baseline_A1C_ge_9_percent" for row in subgroup_rows))

    def test_current_mandros_is_rejected_for_invalid_arm_references(self):
        if not PDF.is_file() or not MANDROS.is_file():
            self.skipTest("Private dapagliflozin inputs are unavailable")
        report = validate_workbook(PDF, MANDROS)
        self.assertEqual(report["recognized_fixture"], "dapagliflozin-monotherapy-dc10-0612-v1")
        self.assertEqual(report["decision"], "REJECT")
        self.assertIn("INVALID_ARM_REFERENCE", {item["code"] for item in report["issues"]})

    def test_corrected_workbook_is_atomic_and_reopenable(self):
        if not OUTPUT.is_file():
            self.skipTest("Corrected workbook has not been generated")
        workbook = load_workbook(OUTPUT, read_only=True, data_only=False)
        self.assertEqual(len(workbook.sheetnames), 14)
        self.assertNotIn("p_value", [cell.value for cell in workbook["Outcomes"][1]])
        test_headers = [cell.value for cell in workbook["Statistical Tests"][1]]
        self.assertIn("p <", test_headers)
        self.assertIn("p =", test_headers)
        self.assertIn("p >", test_headers)
        inventory = list(workbook["Document Inventory"].iter_rows(values_only=True))
        self.assertFalse(any("manual" in str(value).lower() for row in inventory for value in row if value is not None))
        workbook.close()


if __name__ == "__main__":
    unittest.main()
