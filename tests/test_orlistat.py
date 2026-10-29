import json
import unittest
from pathlib import Path

from app.validator import validate_workbook


ROOT = Path(__file__).resolve().parents[1]
DATASET_PATH = ROOT / "standards" / "orlistat_workbook_dataset.json"
PDF_PATH = ROOT / "data" / "private" / "orlistat" / "source.pdf"
MANDROS_PATH = ROOT / "data" / "private" / "orlistat" / "mandros_input.xlsx"
OUTPUT_PATH = ROOT / "outputs" / "orlistat-rct-2024" / "orlistat_Mandros_Blind_corrected.xlsx"


class OrlistatFixtureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.bundle = json.loads(DATASET_PATH.read_text(encoding="utf-8"))

    def test_truth_bundle_covers_all_main_article_tables(self):
        self.assertEqual(len(self.bundle["baseline_pairs"]), 34)
        self.assertEqual(len(self.bundle["outcome_pairs"]), 21)
        self.assertEqual(self.bundle["arms"][0]["allocated_n"], 35)
        self.assertEqual(self.bundle["arms"][1]["allocated_n"], 37)
        self.assertEqual(self.bundle["arms"][0]["analyzed_n"], 28)
        self.assertEqual(self.bundle["arms"][1]["analyzed_n"], 35)
        self.assertEqual(len(self.bundle["document_inventory"]), 1)
        self.assertEqual(self.bundle["document_inventory"][0]["document_type"], "main_article")
        self.assertNotIn("supplement", json.dumps(self.bundle["outcomes"]).lower())

    def test_key_source_values_are_preserved(self):
        bmi = next(row for row in self.bundle["eligibility"] if row["criterion_id"] == "INC_BMI")
        self.assertEqual(bmi["operator_gt"], True)
        self.assertEqual(bmi["operator_eq"], True)
        self.assertEqual(bmi["value"], 25)
        flares = [row for row in self.bundle["outcomes"] if row["outcome"] == "participants with at least one gout flare"]
        observed = {(row["arm_id"], row["event_count"], row["denominator"]) for row in flares}
        self.assertEqual(observed, {("ARM_PBO", 11, 28), ("ARM_ORL", 5, 35)})

    def test_mandros_fixture_is_rejected(self):
        if not PDF_PATH.exists() or not MANDROS_PATH.exists():
            self.skipTest("Private Orlistat fixtures are not installed")
        report = validate_workbook(PDF_PATH, MANDROS_PATH)
        self.assertEqual(report["recognized_fixture"], "orlistat-rct-2024-v1")
        self.assertEqual(report["decision"], "REJECT")
        self.assertEqual(report["score"], 29)
        issue_codes = {issue["code"] for issue in report["issues"]}
        self.assertTrue({"WRONG_BMI_THRESHOLD", "INCOMPLETE_SECONDARY_EFFICACY", "WRONG_OR_MISSING_GOUT_FLARES", "MISSING_ORLISTAT_ADVERSE_EVENTS"}.issubset(issue_codes))

    def test_source_conflict_is_machine_readable(self):
        conflict = self.bundle["source_conflicts"][0]
        self.assertEqual(conflict["published_total_value"], 49)
        self.assertEqual(conflict["reconciled_arm_sum"], 48)
        pair = next(row for row in self.bundle["outcome_pairs"] if row["id"] == "EFF_GOUT_PP")
        self.assertEqual(pair["conflict_id"], conflict["conflict_id"])
        conflict_ids = {row["conflict_id"] for row in self.bundle["source_conflicts"]}
        self.assertIn("CONFLICT_TABLE1_GOUT_HISTORY", conflict_ids)

    def test_unprinted_recurrent_event_counts_are_not_created(self):
        recurrent_rows = [row for row in self.bundle["outcomes"] if row["outcome"] == "participants with recurrent gout flare"]
        self.assertEqual(recurrent_rows, [])
        recurrent_test = next(row for row in self.bundle["statistical_tests"] if row["test_id"] == "TEST_RECURRENT_GOUT_HR")
        self.assertEqual(recurrent_test["effect_value"], 0.27)
        self.assertEqual(recurrent_test["ci_lower"], 0.10)
        self.assertEqual(recurrent_test["ci_upper"], 0.75)

    def test_exported_workbook_has_phase2_metadata_without_supplementary_results(self):
        if not OUTPUT_PATH.exists():
            self.skipTest("Corrected Orlistat workbook has not been built")
        from openpyxl import load_workbook
        workbook = load_workbook(OUTPUT_PATH, data_only=False, read_only=True)
        self.assertIn("Document Inventory", workbook.sheetnames)
        self.assertIn("Source Conflicts", workbook.sheetnames)
        self.assertIn("Phase 2 Readiness", workbook.sheetnames)
        self.assertIn("Statistical Methods", workbook.sheetnames)
        arm_headers = [cell.value for cell in next(workbook["Arms"].iter_rows(min_row=1, max_row=1))]
        self.assertTrue({"n_randomized", "n_baseline", "n_analyzed"}.issubset(set(arm_headers)))
        outcome_headers = [cell.value for cell in next(workbook["Outcomes"].iter_rows(min_row=1, max_row=1))]
        self.assertTrue({"estimate_type", "source_status", "conflict_id", "canonical_record_id", "independent_study_id", "estimand_family", "meta_analysis_eligibility", "meta_analysis_reason", "include_in_meta_analysis"}.issubset(set(outcome_headers)))
        workbook_text = " ".join(
            str(value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if value is not None
        ).lower()
        self.assertNotIn("supplement", workbook_text)


if __name__ == "__main__":
    unittest.main()
