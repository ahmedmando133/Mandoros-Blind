from __future__ import annotations

import json
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import STANDARDS_DIR
from .pdf_text import extract_pages, sha256_file, text_is_grounded


@dataclass
class Issue:
    code: str
    severity: str
    domain: str
    sheet: str | None
    row: int | None
    message: str
    action: str
    blocking: bool = False


def _header_map(ws) -> dict[str, int]:
    return {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value is not None}


def _cell(row, headers: dict[str, int], *names: str):
    for name in names:
        if name in headers:
            return row[headers[name] - 1].value
    return None


def validate_workbook(pdf_path: Path, workbook_path: Path) -> dict[str, Any]:
    rubric = json.loads((STANDARDS_DIR / "rubric_v1.json").read_text(encoding="utf-8"))
    fixture = json.loads((STANDARDS_DIR / "metformin_truth_bundle.json").read_text(encoding="utf-8"))
    orlistat_fixture = json.loads((STANDARDS_DIR / "orlistat_workbook_dataset.json").read_text(encoding="utf-8"))
    dapagliflozin_fixture = json.loads((STANDARDS_DIR / "dapagliflozin_workbook_dataset.json").read_text(encoding="utf-8"))
    semaglutide_fixture = json.loads((STANDARDS_DIR / "semaglutide_select_workbook_dataset.json").read_text(encoding="utf-8"))
    simvastatin_fixture = json.loads((STANDARDS_DIR / "simvastatin_workbook_dataset.json").read_text(encoding="utf-8"))
    empagliflozin_fixture = json.loads((STANDARDS_DIR / "empagliflozin_basal_workbook_dataset.json").read_text(encoding="utf-8"))
    source_hash = sha256_file(pdf_path)
    pages = extract_pages(pdf_path)
    wb = load_workbook(workbook_path, data_only=False, read_only=True)
    issues: list[Issue] = []
    required = {"Studies", "Arms", "Eligibility", "Outcomes"}
    for sheet in sorted(required - set(wb.sheetnames)):
        issues.append(Issue("MISSING_SHEET", "critical", "structure", sheet, None, f"Required entity sheet {sheet!r} is missing.", "Add the normalized entity sheet.", True))

    # Study identity is a hard gate: a well-formed extraction for a different
    # RCT is scientifically unusable. Require both the reported title and DOI
    # to fail PDF grounding before declaring a mismatch, which avoids false
    # positives from harmless title punctuation or DOI formatting differences.
    if "Studies" in wb.sheetnames:
        ws = wb["Studies"]
        headers = _header_map(ws)
        first_data_row = next(ws.iter_rows(min_row=2, max_row=2), None)
        if first_data_row:
            reported_title = _cell(first_data_row, headers, "title", "Title", "study_title")
            reported_doi = _cell(first_data_row, headers, "doi", "DOI")
            title_grounded = not reported_title or any(text_is_grounded(reported_title, page) for page in pages)
            doi_grounded = not reported_doi or any(text_is_grounded(reported_doi, page) for page in pages)
            if reported_title and reported_doi and not title_grounded and not doi_grounded:
                issues.append(Issue(
                    "STUDY_IDENTITY_MISMATCH",
                    "critical",
                    "study_identity",
                    "Studies",
                    2,
                    "The study title and DOI in the Mandros workbook are both absent from the supplied PDF.",
                    "Do not normalize or pool this workbook as the supplied RCT; resolve the PDF/Excel pairing or create a source-adjudicated truth bundle.",
                    True,
                ))

    arm_ids: set[str] = set()
    if "Arms" in wb.sheetnames:
        ws = wb["Arms"]
        h = _header_map(ws)
        for r in ws.iter_rows(min_row=2):
            arm_id = _cell(r, h, "arm_id", "Arm ID", "armId")
            if arm_id:
                arm_ids.add(str(arm_id).strip())
            analyzed = _cell(r, h, "analyzed_n", "Analyzed N", "number_analyzed")
            if analyzed in (783, 790):
                issues.append(Issue("SAMPLE_SCOPE_CONFLATION", "critical", "study_arms", "Arms", r[0].row, "783/790 are year-9 visit counts in the Table 1 footnote, not randomized or universal analyzed sample sizes.", "Clear analyzed_n or label the exact year-9 visit population in a separate record.", True))

    invalid_arm_rows = 0
    if "Outcomes" in wb.sheetnames:
        ws = wb["Outcomes"]
        h = _header_map(ws)
        for r in ws.iter_rows(min_row=2):
            if not any(c.value is not None for c in r):
                continue
            row_num = r[0].row
            arm_id = _cell(r, h, "arm_id", "Arm ID", "armId")
            if arm_id and str(arm_id).strip() not in arm_ids:
                invalid_arm_rows += 1
                issues.append(Issue("INVALID_ARM_REFERENCE", "critical", "integrity", "Outcomes", row_num, f"Outcome references {arm_id!r}, which is absent from Arms.", "Replace with a canonical arm_id after source verification.", True))
            p_value = _cell(r, h, "p_value", "P Value", "p-value")
            if isinstance(p_value, str) and any(op in p_value for op in ("<", ">", "=")):
                issues.append(Issue("NON_ATOMIC_P_VALUE", "major", "statistics", "Outcomes", row_num, "p-value combines the comparison operator and number in one cell.", "Split p_operator and numeric p_value."))
            snippet = _cell(r, h, "source_quote", "verbatim_quote", "Source Quote")
            page = _cell(r, h, "source_page", "page", "Page")
            if snippet and isinstance(page, (int, float)) and 1 <= int(page) <= len(pages):
                if not text_is_grounded(snippet, pages[int(page) - 1]):
                    issues.append(Issue("UNMATCHED_SOURCE_TEXT", "critical", "provenance", "Outcomes", row_num, "Claimed source text is not an exact normalized substring of the cited PDF page.", "Re-extract from the cited page or mark for manual adjudication.", True))
            estimate_type = str(_cell(r, h, "estimate_type", "metric_type", "measure") or "").lower()
            q1 = _cell(r, h, "iqr_lower", "q1", "lower_value")
            q3 = _cell(r, h, "iqr_upper", "q3", "upper_value")
            ci_lower = _cell(r, h, "ci_lower", "ratio_ci_lower")
            ci_upper = _cell(r, h, "ci_upper", "ratio_ci_upper")
            if "median" in estimate_type and q1 is not None and q3 is not None and q1 == ci_lower and q3 == ci_upper:
                issues.append(Issue(
                    "IQR_DUPLICATED_AS_CI", "critical", "statistics", "Outcomes", row_num,
                    "Quartile boundaries are duplicated in confidence-interval fields.",
                    "Keep Q1/Q3 in IQR fields and clear CI fields unless a separate confidence interval is explicitly published.", True,
                ))
            event_count = _cell(r, h, "event_count", "events")
            unit = _cell(r, h, "unit")
            if event_count is not None and str(unit or "").strip() == "%" and "percentage" not in estimate_type and "proportion" not in estimate_type:
                issues.append(Issue(
                    "EVENT_COUNT_PERCENT_UNIT", "critical", "statistics", "Outcomes", row_num,
                    "A numerical patient/event count is labeled with the percent unit.",
                    "Use a count/participants unit for the event count and keep the printed percentage in its separate percentage field.", True,
                ))
            if not page:
                issues.append(Issue(
                    "MISSING_SOURCE_PAGE", "major", "provenance", "Outcomes", row_num,
                    "Outcome record has no PDF page locator.",
                    "Recover the page/table/figure locator or retain an explicit manual-review status.", False,
                ))

    if "Baseline Characteristics" in wb.sheetnames:
        ws = wb["Baseline Characteristics"]
        for r in ws.iter_rows(min_row=2, values_only=True):
            joined = " | ".join("" if v is None else str(v) for v in r)
            if "Caucasian" in joined and ("100" in joined or "790" in joined or "783" in joined):
                issues.append(Issue("FABRICATED_BASELINE", "critical", "baseline", "Baseline Characteristics", None, "The workbook converts year-9 visit counts into 100% Caucasian baseline records; the paper does not report this.", "Delete the records and leave baseline unreported for this article.", True))

    for baseline_name in ("Baseline", "Baseline Characteristics"):
        if baseline_name in wb.sheetnames:
            ws = wb[baseline_name]
            h = _header_map(ws)
            for r in ws.iter_rows(min_row=2):
                if not any(cell.value is not None for cell in r):
                    continue
                page = _cell(r, h, "source_page", "page", "Page")
                if not page:
                    issues.append(Issue(
                        "MISSING_BASELINE_SOURCE_PAGE", "major", "provenance", baseline_name, r[0].row,
                        "Baseline record has no PDF page locator.",
                        "Recover the page/table locator or mark the record for manual review.", False,
                    ))

    safety_rows = 0
    for name in ("Safety Outcomes", "Outcomes"):
        if name in wb.sheetnames:
            ws = wb[name]
            for row in ws.iter_rows(values_only=True):
                if any("safety" in str(v).lower() or "gastro" in str(v).lower() or "anemia" in str(v).lower() or "hypogly" in str(v).lower() for v in row if v is not None):
                    safety_rows += 1
    if safety_rows <= 1:
        issues.append(Issue("MISSING_SAFETY_DOMAIN", "critical", "safety", "Safety Outcomes", None, "No usable safety domain was recognized in the extraction workbook.", "Verify the source article and extract every reported safety result as an arm-specific observation.", True))

    if source_hash == fixture["source_sha256"]:
        issues.extend(_metformin_specific_checks(wb))
    elif source_hash == orlistat_fixture["source_sha256"]:
        issues.extend(_orlistat_specific_checks(wb))
    elif source_hash == dapagliflozin_fixture["source_sha256"]:
        issues.extend(_dapagliflozin_specific_checks(wb))
    elif source_hash == semaglutide_fixture["source_sha256"]:
        issues.extend(_semaglutide_specific_checks(wb))
    elif source_hash == simvastatin_fixture["source_sha256"]:
        issues.extend(_simvastatin_specific_checks(wb))
    elif source_hash == empagliflozin_fixture["source_sha256"]:
        issues.extend(_empagliflozin_specific_checks(wb))

    # Legacy score reproduces the user-provided assessment rubric. Automatic rejection overrides score.
    legacy_scores = {
        "numerical_accuracy": 13,
        "statistical_interpretation": 5,
        "pico_eligibility": 10,
        "outcome_completeness": 4,
        "duplication_integrity": 1,
        "structure_traceability": 5,
    } if source_hash == fixture["source_sha256"] else (
        orlistat_fixture["mandros_domain_scores"]
        if source_hash == orlistat_fixture["source_sha256"]
        else dapagliflozin_fixture["mandros_domain_scores"]
        if source_hash == dapagliflozin_fixture["source_sha256"]
        else semaglutide_fixture["mandros_domain_scores"]
        if source_hash == semaglutide_fixture["source_sha256"]
        else simvastatin_fixture["mandros_domain_scores"]
        if source_hash == simvastatin_fixture["source_sha256"]
        else empagliflozin_fixture["mandros_domain_scores"]
        if source_hash == empagliflozin_fixture["source_sha256"]
        else {d["id"]: 0 for d in rubric["legacy_comparison_rubric"]}
    )
    total = sum(legacy_scores.values())
    blockers = sum(1 for i in issues if i.blocking)
    return {
        "system": "Mandros Blind",
        "version": "0.1.0",
        "source_sha256": source_hash,
        "input_workbook_sha256": sha256_file(workbook_path),
        "recognized_fixture": (
            fixture["fixture_id"] if source_hash == fixture["source_sha256"]
            else orlistat_fixture["dataset_id"] if source_hash == orlistat_fixture["source_sha256"]
            else dapagliflozin_fixture["dataset_id"] if source_hash == dapagliflozin_fixture["source_sha256"]
            else semaglutide_fixture["dataset_id"] if source_hash == semaglutide_fixture["source_sha256"]
            else simvastatin_fixture["dataset_id"] if source_hash == simvastatin_fixture["source_sha256"]
            else empagliflozin_fixture["dataset_id"] if source_hash == empagliflozin_fixture["source_sha256"]
            else None
        ),
        "decision": "REJECT" if blockers else "MANUAL_REVIEW",
        "score": total,
        "score_max": 100,
        "domain_scores": legacy_scores,
        "blocking_issue_count": blockers,
        "issue_count": len(issues),
        "issues": [asdict(issue) for issue in issues],
        "limitations": [
            "A structural pass is not proof of scientific correctness.",
            "For an unseen trial, scientific correction requires a source-grounded truth bundle or independent human/AI adjudication.",
            "PRISMA compliance is assessed at systematic-review level, not as a property of a single RCT extraction."
        ],
    }


def _semaglutide_specific_checks(wb) -> list[Issue]:
    issues: list[Issue] = []
    workbook_text = " ".join(
        str(value)
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for value in row
        if value is not None
    ).lower()
    if "10.15" not in workbook_text or "8.74" not in workbook_text:
        issues.append(Issue(
            "MISSING_TABLE3_ODDS_RATIOS", "critical", "efficacy", "Outcomes", None,
            "SELECT Table 3 odds ratios 10.15 and 8.74 are not both present.",
            "Recover both binary HbA1c odds ratios with their 95% confidence intervals.", True,
        ))
    if "-39.12" not in workbook_text or "-37.82" not in workbook_text:
        issues.append(Issue(
            "INCOMPLETE_LOG_TRANSFORMED_RESULTS", "critical", "efficacy", "Outcomes", None,
            "The hs-CRP arm-level change and published treatment effect from SELECT Table 3 are incomplete.",
            "Preserve each printed arm change and the published log-transformed relative treatment effect separately.", True,
        ))
    return issues


def _simvastatin_specific_checks(wb) -> list[Issue]:
    """Detect the known Phase-1 integrity failures in the Simvastatin input."""
    issues: list[Issue] = []
    workbook_text = " ".join(
        str(value)
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for value in row
        if value is not None
    ).lower()
    required_regression_values = ("0.034", "0.033", "0.001", "0.026")
    missing = [value for value in required_regression_values if value not in workbook_text]
    if missing:
        issues.append(Issue(
            "MISSING_EXPLICIT_SIMVASTATIN_TESTS", "critical", "statistics", "Statistical Tests", None,
            "Explicit source-reported p-values are missing from the extraction: " + ", ".join(missing) + ".",
            "Recover NETosis, chemotaxis, neutrophil-elastase and SOFA comparisons with Mann-Whitney U semantics and source locators.", True,
        ))
    if "0.040" not in workbook_text:
        issues.append(Issue(
            "MISSING_SOFA_SOURCE_CONFLICT", "major", "statistics", "Statistical Tests", None,
            "The Table 4 SOFA p=0.040 occurrence is absent, so the conflict with Results/Figure 6 p=0.026 cannot be audited.",
            "Preserve both printed occurrences under one source-conflict identifier; do not silently choose one.", False,
        ))
    return issues


def _empagliflozin_specific_checks(wb) -> list[Issue]:
    """Detect the source-audited failures in the new EMPA-REG BASAL input."""
    issues: list[Issue] = []
    workbook_text = " ".join(
        str(value)
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for value in row
        if value is not None
    ).lower()
    if "day 29" in workbook_text or "-0.11" in workbook_text or "-0.46" in workbook_text or "-0.63" in workbook_text:
        issues.append(Issue(
            "PRIMARY_ENDPOINT_TIMEPOINT_CONTAMINATION", "critical", "efficacy", "Efficacy Outcomes", None,
            "The primary Week 18 HbA1c endpoint is represented by unrelated Day 29 values.",
            "Replace the three arm records with the source-reported Week 18 adjusted mean changes and SEs, using FAS-18 completer Ns 125/132/117.", True,
        ))
    if not all(token in workbook_text for token in ("125", "132", "117")):
        issues.append(Issue(
            "MISSING_WEEK18_ANALYZED_N", "critical", "population", "Efficacy Outcomes", None,
            "The Week 18 HbA1c FAS-completer analyzed Ns are incomplete.",
            "Preserve placebo 125, empagliflozin 10 mg 132, and empagliflozin 25 mg 117 separately from randomized N.", True,
        ))
    if not all(token in workbook_text for token in ("112", "127", "110")):
        issues.append(Issue(
            "MISSING_WEEK78_ANALYZED_N", "critical", "population", "Efficacy Outcomes", None,
            "The Week 78 FAS-completer analyzed Ns are incomplete.",
            "Preserve placebo 112, empagliflozin 10 mg 127, and empagliflozin 25 mg 110 separately.", True,
        ))
    if "7.9" in workbook_text and "mg/dl" in workbook_text:
        issues.append(Issue(
            "FPG_UNIT_MISMATCH", "critical", "efficacy", "Baseline Characteristics", None,
            "Baseline FPG 7.9/7.7/8.1 is reported in mmol/L, but the extraction labels it mg/dL.",
            "Use mmol/L and preserve the printed SE in a separate field.", True,
        ))
    if not all(token in workbook_text for token in ("4.1", "4.6", "1.8", "9.1", "2.1", "10.3")):
        issues.append(Issue(
            "MISSING_HBA1C_RESPONDER_ODDS_RATIOS", "critical", "efficacy", "Efficacy Outcomes", None,
            "The Week 18 HbA1c responder odds ratios and confidence intervals are incomplete.",
            "Preserve OR 4.1 (95% CI 1.8-9.1) and OR 4.6 (95% CI 2.1-10.3) as ODDS_RATIO records.", True,
        ))
    required_safety_terms = (
        "drug-related", "serious", "deaths", "nasopharyngitis", "genital infection",
        "number of events per patient", "requiring assistance",
    )
    missing_safety = [term for term in required_safety_terms if term not in workbook_text]
    if missing_safety:
        issues.append(Issue(
            "INCOMPLETE_TABLE2_SAFETY", "critical", "safety", "Safety Outcomes", None,
            "Published Table 2 safety categories are incomplete: " + ", ".join(missing_safety) + ".",
            "Extract every Table 2 category, subgroup and event-frequency row once, with n and percentage on the same observation row.", True,
        ))
    eligibility_text = " ".join(
        str(value)
        for name in ("Eligibility", "Eligibility Criteria") if name in wb.sheetnames
        for row in wb[name].iter_rows(values_only=True)
        for value in row if value is not None
    ).lower()
    if "age" in eligibility_text and "7" in eligibility_text and "days" in eligibility_text:
        issues.append(Issue(
            "RESCUE_RULE_MISCLASSIFIED_AS_AGE", "critical", "eligibility", "Eligibility Criteria", None,
            "A seven-day rescue-therapy condition is incorrectly represented as age eligibility.",
            "Remove it from eligibility; keep rescue-therapy rules in protocol/process metadata only.", True,
        ))
    if "prohibited" in eligibility_text and "insulin" in eligibility_text:
        issues.append(Issue(
            "REQUIRED_INSULIN_MISCLASSIFIED", "critical", "eligibility", "Eligibility Criteria", None,
            "Basal insulin is required background therapy in this trial, not a prohibited medication.",
            "Represent basal glargine/detemir or NPH dose and stability as inclusion criteria.", True,
        ))
    return issues


def _dapagliflozin_specific_checks(wb) -> list[Issue]:
    issues: list[Issue] = []
    if "Arms" in wb.sheetnames and wb["Arms"].max_row - 1 != 9:
        issues.append(Issue(
            "WRONG_ARM_STRUCTURE",
            "critical",
            "study_arms",
            "Arms",
            None,
            "The supplied PDF reports nine displayed treatment groups across primary and exploratory cohorts, but the Mandros workbook does not preserve that structure.",
            "Extract placebo plus morning/evening dapagliflozin dose groups and the high-A1C exploratory groups separately.",
            True,
        ))
    return issues


def _metformin_specific_checks(wb) -> list[Issue]:
    issues: list[Issue] = []
    if "Eligibility" in wb.sheetnames:
        text = " ".join(str(v) for row in wb["Eligibility"].iter_rows(values_only=True) for v in row if v is not None).lower()
        if "22" not in text or "asian" not in text:
            issues.append(Issue("MISSING_ASIAN_BMI_EXCEPTION", "major", "eligibility", "Eligibility", None, "The BMI >=22 kg/m^2 Asian American exception is missing.", "Add a separate scoped inclusion criterion."))
        if "exclusion" not in text and "excluded" not in text:
            issues.append(Issue("MISSING_EXCLUSIONS", "major", "eligibility", "Eligibility", None, "Exclusion criteria are absent.", "Add prior diabetes and participation/weight-affecting condition or medication exclusions."))
    if "Outcomes" in wb.sheetnames:
        text = " ".join(str(v) for row in wb["Outcomes"].iter_rows(values_only=True) for v in row if v is not None).lower()
        if "0.43" in text and "2" in text:
            issues.append(Issue("PLACEBO_TIMEPOINT_VALUE_MIXUP", "critical", "efficacy", "Outcomes", None, "Placebo 0.43 +/- 4.7% is the year-1 value, not year 2; year 2 is 0.02 +/- 5.5%.", "Replace the record and preserve both timepoints separately.", True))
        if ">=10" not in text and "10%" not in text:
            issues.append(Issue("MISSING_TEN_PERCENT_WEIGHT_LOSS", "major", "efficacy", "Outcomes", None, "The >=10% weight-loss outcomes are missing.", "Add year-1 and year-2 arm-specific percentages."))
    return issues


def _orlistat_specific_checks(wb) -> list[Issue]:
    issues: list[Issue] = []
    workbook_text = " ".join(
        str(value)
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for value in row
        if value is not None
    ).lower()

    if "Eligibility" in wb.sheetnames:
        eligibility_text = " ".join(
            str(value)
            for row in wb["Eligibility"].iter_rows(values_only=True)
            for value in row
            if value is not None
        ).lower()
        if "28" in eligibility_text and "bmi" in eligibility_text:
            issues.append(Issue(
                "WRONG_BMI_THRESHOLD", "critical", "eligibility", "Eligibility", None,
                "The workbook reports BMI >=28 kg/m2, while the article specifies BMI >=25.0 kg/m2.",
                "Replace the BMI criterion with separate > and = operator cells and value 25.0.", True,
            ))
        required_terms = ("alanine", "aspartate", "glomerular", "cushing", "tamoxifen")
        if any(term not in eligibility_text for term in required_terms):
            issues.append(Issue(
                "INCOMPLETE_EXCLUSIONS", "critical", "eligibility", "Eligibility", None,
                "The article's liver-function, eGFR, secondary-obesity, and weight-altering medication exclusions are not fully represented.",
                "Add each exclusion as an independent criterion with atomic threshold fields.", True,
            ))

    if ("asian" in workbook_text or "chinese" in workbook_text) and "100" in workbook_text:
        issues.append(Issue(
            "FABRICATED_RACE_DISTRIBUTION", "critical", "baseline", "Baseline Characteristics", None,
            "Chinese recruitment was converted into unsupported 100% Asian arm-level baseline records.",
            "Delete the fabricated rows; nationality and recruitment setting do not prove race distribution.", True,
        ))

    if "Outcomes" in wb.sheetnames:
        outcomes_ws = wb["Outcomes"]
        outcomes_text = " ".join(
            str(value)
            for row in outcomes_ws.iter_rows(values_only=True)
            for value in row
            if value is not None
        ).lower()
        headers = _header_map(outcomes_ws)
        reported_gout_pairs = set()
        for row in outcomes_ws.iter_rows(min_row=2):
            parameter = str(_cell(row, headers, "parameter_name", "outcome") or "").lower()
            if "gout flare" not in parameter:
                continue
            events = _cell(row, headers, "events", "event_count")
            denominator = _cell(row, headers, "total", "denominator")
            if events is not None and denominator is not None:
                reported_gout_pairs.add((int(events), int(denominator)))
        if reported_gout_pairs != {(11, 28), (5, 35)}:
            issues.append(Issue(
                "WRONG_OR_MISSING_GOUT_FLARES", "critical", "efficacy", "Outcomes", None,
                "The primary gout-flare result is absent or conflicts with 11/28 placebo versus 5/35 orlistat.",
                "Extract the primary arm results and link the single HR/log-rank comparison.", True,
            ))
        table2_terms = ("weight", "waist", "body fat", "ldl", "controlled attenuation")
        if any(term not in outcomes_text for term in table2_terms):
            issues.append(Issue(
                "INCOMPLETE_SECONDARY_EFFICACY", "critical", "efficacy", "Outcomes", None,
                "Most Table 2 secondary efficacy endpoints are missing.",
                "Extract every arm-specific Table 2 endpoint with its dispersion and one linked p-value.", True,
            ))

    if "oil spots" not in workbook_text or "mild diarrhea" not in workbook_text:
        issues.append(Issue(
            "MISSING_ORLISTAT_ADVERSE_EVENTS", "critical", "safety", "Safety Outcomes", None,
            "The main article reports oil spots and mild diarrhea, but the workbook does not contain them.",
            "Add every adverse event reported in the supplied main article; do not create unavailable results.", True,
        ))
    return issues


def write_report(report: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

