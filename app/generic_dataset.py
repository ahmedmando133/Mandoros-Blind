from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .pdf_text import extract_pages, normalize_text, sha256_file, text_is_grounded


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _records(workbook, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        return []
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration:
        return []
    result: list[dict[str, Any]] = []
    for values in rows:
        if not any(value not in (None, "") for value in values):
            continue
        result.append({header: _json_value(value) for header, value in zip(headers, values) if header})
    return result


def _first(record: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = record.get(key)
        if value not in (None, ""):
            return value
    return None


def _safe_id(value: Any, prefix: str) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", "_", str(value or "")).strip("_").upper()
    return text[:80] or prefix


def _truth(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "y"}


def _source_ids(value: Any) -> list[str]:
    return [item for item in re.split(r"[;,|\s]+", str(value or "").strip()) if item]


def _number_tokens(values: Iterable[Any]) -> list[str]:
    tokens: list[str] = []
    for value in values:
        if isinstance(value, bool) or value in (None, ""):
            continue
        if isinstance(value, (int, float)):
            token = f"{value:g}" if isinstance(value, float) else str(value)
            tokens.append(token)
            continue
        tokens.extend(re.findall(r"(?<![A-Za-z])[-+]?\d+(?:\.\d+)?", str(value).replace(",", "")))
    return list(dict.fromkeys(tokens))


def _number_on_page(token: str, page_text: str) -> bool:
    haystack = normalize_text(page_text).replace(",", "")
    normalized = token.lstrip("+")
    alternatives = {normalized}
    try:
        number = float(normalized)
        alternatives.add(f"{number:g}")
        if number.is_integer():
            alternatives.add(str(int(number)))
    except ValueError:
        pass
    return any(re.search(rf"(?<!\d){re.escape(item)}(?!\d)", haystack) for item in alternatives if item)


def _context_words(*values: Any) -> list[str]:
    stop = {
        "with", "from", "that", "this", "were", "was", "patients", "patient", "group",
        "value", "mean", "change", "baseline", "weeks", "days", "month", "months",
    }
    words: list[str] = []
    for value in values:
        for word in re.findall(r"[a-zA-Z][a-zA-Z-]{3,}", normalize_text(value)):
            if word not in stop:
                words.append(word)
    return list(dict.fromkeys(words))


def _verification_for_record(
    record: dict[str, Any],
    provenance: dict[str, dict[str, Any]],
    pages: list[str],
    label: Any,
) -> tuple[str, int | None, str | None]:
    entries = [provenance[item] for item in _source_ids(record.get("source_ids")) if item in provenance]
    if not entries:
        page = _first(record, "source_page", "page")
        entries = [{"page": page, "evidence": _first(record, "raw_source_text", "source_evidence")}]

    for entry in entries:
        evidence = _first(entry, "evidence", "exact_source_text")
        page = _first(entry, "page", "source_page")
        if isinstance(page, (int, float)) and 1 <= int(page) <= len(pages) and evidence:
            if text_is_grounded(evidence, pages[int(page) - 1]):
                return "published_main_article_exact", int(page), str(evidence)

    for entry in entries:
        evidence = _first(entry, "evidence", "exact_source_text")
        page = _first(entry, "page", "source_page")
        if not isinstance(page, (int, float)) or not (1 <= int(page) <= len(pages)):
            continue
        page_text = pages[int(page) - 1]
        numeric_values = [
            value for key, value in record.items()
            if key not in {"source_ids", "source_page", "page"} and isinstance(value, (int, float)) and not isinstance(value, bool)
        ]
        numbers = _number_tokens(numeric_values)
        words = _context_words(label, evidence)
        page_normalized = normalize_text(page_text)
        numbers_match = bool(numbers) and all(_number_on_page(token, page_text) for token in numbers)
        context_match = any(word in page_normalized for word in words)
        if numbers_match and context_match:
            return "published_main_article_numeric_context", int(page), str(evidence or "") or None

    for page_number, page_text in enumerate(pages, start=1):
        evidence = _first(record, "raw_source_text", "source_evidence")
        if evidence and text_is_grounded(evidence, page_text):
            return "published_main_article_exact", page_number, str(evidence)
    page = _first(record, "source_page", "page")
    return "requires_manual_source_review", int(page) if isinstance(page, (int, float)) else None, None


def _operator_flags(*values: Any) -> tuple[bool, bool, bool]:
    text = "".join(str(value or "") for value in values)
    return "<" in text, "=" in text, ">" in text


def _parse_p_value(value: Any, operator: Any) -> tuple[str | None, float | None]:
    op_text = str(operator or "").strip()
    if value in (None, ""):
        return (op_text if op_text in {"<", "=", ">"} else None), None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return (op_text if op_text in {"<", "=", ">"} else "="), float(value)
    text = str(value).strip()
    match = re.search(r"([<=>])?\s*(\d+(?:\.\d+)?)", text)
    if not match:
        return None, None
    return match.group(1) or (op_text if op_text in {"<", "=", ">"} else "="), float(match.group(2))


def _study_record(rows: list[dict[str, Any]], pdf_path: Path, pages: list[str]) -> dict[str, Any]:
    source = rows[0] if rows else {}
    study_id = _safe_id(_first(source, "study_id") or pdf_path.stem, "STUDY")
    title = _first(source, "title") or pdf_path.stem
    title_page = 1
    for index, page in enumerate(pages, start=1):
        if text_is_grounded(title, page):
            title_page = index
            break
    return {
        "study_id": study_id,
        "title": title,
        "authors": _first(source, "authors"),
        "year": _first(source, "publication_year", "year"),
        "journal": _first(source, "journal"),
        "doi": _first(source, "doi"),
        "design": _first(source, "study_design", "design"),
        "study_phase": _first(source, "study_phase"),
        "randomized_total_n": _first(source, "total_sample_size", "randomized_total_n"),
        "duration_value": _first(source, "duration_value"),
        "duration_unit": _first(source, "duration_unit"),
        "condition": _first(source, "condition_studied", "condition"),
        "source_page": title_page,
        "source_status": "published_main_article_exact" if text_is_grounded(title, pages[title_page - 1]) else "requires_manual_source_review",
    }


def _arms(
    rows: list[dict[str, Any]], provenance: dict[str, dict[str, Any]], pages: list[str]
) -> list[dict[str, Any]]:
    result = []
    for index, source in enumerate(rows, start=1):
        label = _first(source, "name", "arm_name") or f"Arm {index}"
        status, page, _ = _verification_for_record(source, provenance, pages, label)
        result.append({
            "arm_id": _safe_id(_first(source, "arm_id"), f"ARM_{index:03d}"),
            "label": label,
            "n_randomized": _first(source, "randomized_n", "sample_size_allocated"),
            "n_baseline": None,
            "n_analyzed": _first(source, "analyzed_n", "sample_size_analyzed"),
            "analysis_population": "as_reported",
            "drug": _first(source, "drug", "drug_name"),
            "dose_value": _first(source, "dose_value"),
            "dose_unit": _first(source, "dose_unit"),
            "initial_frequency_value": _first(source, "frequency_value", "frequency_number"),
            "initial_frequency_unit": _first(source, "frequency_per_unit", "frequency_unit"),
            "target_frequency_value": _first(source, "frequency_value", "frequency_number"),
            "target_frequency_unit": _first(source, "frequency_per_unit", "frequency_unit"),
            "administration_timing": None,
            "duration_value": _first(source, "duration_value"),
            "duration_unit": _first(source, "duration_unit"),
            "titration_value": None,
            "titration_unit": None,
            "route": _first(source, "route"),
            "source_page": page,
            "source_status": status,
        })
    return result


def _eligibility(
    rows: list[dict[str, Any]], provenance: dict[str, dict[str, Any]], pages: list[str]
) -> list[dict[str, Any]]:
    result = []
    for index, source in enumerate(rows, start=1):
        lower = _first(source, "lower_value", "min_value")
        upper = _first(source, "upper_value", "max_value")
        lower_lt, lower_eq, lower_gt = _operator_flags(
            _first(source, "lower_operator_1", "min_operator (>=)"), source.get("lower_operator_2")
        )
        upper_lt, upper_eq, upper_gt = _operator_flags(
            _first(source, "upper_operator_1", "max_operator (<=)"), source.get("upper_operator_2")
        )
        description = _first(source, "description", "raw_source_text")
        variable = _first(source, "variable", "criterion_name") or f"criterion_{index}"
        status, page, _ = _verification_for_record(source, provenance, pages, variable)
        record = {
            "criterion_id": _safe_id(_first(source, "criterion_id"), f"ELIG_{index:03d}"),
            "category": str(_first(source, "kind", "criterion_type") or "unspecified").lower(),
            "variable": variable,
            "operator_lt": False,
            "operator_eq": False,
            "operator_gt": False,
            "value": None,
            "unit": _first(source, "unit"),
            "range_min": None,
            "range_max": None,
            "range_unit": None,
            "range_min_inclusive": None,
            "range_max_inclusive": None,
            "population_scope": _first(source, "population_subgroup", "applicable_population") or "overall",
            "qualifier": description or _first(source, "medication"),
            "source_page": page,
            "source_status": status,
        }
        if lower is not None and upper is not None:
            record.update({
                "range_min": lower,
                "range_max": upper,
                "range_unit": _first(source, "unit"),
                "range_min_inclusive": lower_eq,
                "range_max_inclusive": upper_eq,
            })
        elif lower is not None:
            record.update({"value": lower, "operator_lt": lower_lt, "operator_eq": lower_eq, "operator_gt": lower_gt})
        elif upper is not None:
            record.update({"value": upper, "operator_lt": upper_lt, "operator_eq": upper_eq, "operator_gt": upper_gt})
        result.append(record)
    return result


def _analysis_status(source_status: str, ready: bool, reason: str) -> tuple[str, str, bool]:
    if source_status != "published_main_article_exact":
        return "REQUIRES_REVIEW", "PDF evidence was not an exact page-level text match; verify this record before pooling.", False
    if ready:
        return "READY", reason, True
    return "REQUIRES_REVIEW", reason, False


def _split_characteristic_and_subgroup(value: Any, explicit_subgroup: Any = None) -> tuple[str, str | None, str | None]:
    """Keep a baseline characteristic and its category in separate fields.

    Mandros commonly emits labels such as ``Sex: Male [Male]`` and
    ``Race: White [White]``.  The normalized workbook must instead store
    ``sex``/``race`` as the characteristic and the category in
    ``subgroup_value`` so a computer can group the observations reliably.
    """
    text = str(value or "").strip()
    supplied = str(explicit_subgroup).strip() if explicit_subgroup not in (None, "") else None
    match = re.match(r"^(sex|gender|race|ethnicity)\s*:\s*([^\[]+?)(?:\s*\[([^\]]+)\])?$", text, re.I)
    if not match:
        return text, ("reported_subgroup" if supplied else None), supplied
    characteristic = "sex" if match.group(1).lower() in {"sex", "gender"} else match.group(1).lower()
    subgroup = supplied or match.group(3) or match.group(2)
    return characteristic, characteristic, str(subgroup).strip().lower()


def _split_outcome_threshold(value: Any) -> tuple[str, bool, bool, bool, Any, str | None]:
    """Separate a published threshold from an outcome label without inference."""
    text = str(value or "").strip().replace("≤", "<=").replace("≥", ">=")
    match = re.match(r"^(.*?)\s*(<=|>=|<|>|=)\s*(-?\d+(?:\.\d+)?)\s*([^\s]*)\s*$", text)
    if not match:
        return text, False, False, False, None, None
    base, operator, number, unit = match.groups()
    return (
        base.strip(),
        "<" in operator,
        "=" in operator,
        ">" in operator,
        float(number),
        unit or None,
    )


def _baseline(
    rows: list[dict[str, Any]], provenance: dict[str, dict[str, Any]], pages: list[str], study_id: str
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for index, source in enumerate(rows, start=1):
        raw_characteristic = _first(source, "variable", "characteristic") or f"baseline_{index}"
        characteristic, subgroup_type, subgroup_value = _split_characteristic_and_subgroup(
            raw_characteristic, _first(source, "population", "gender", "race")
        )
        status, page, _ = _verification_for_record(source, provenance, pages, characteristic)
        base_id = _safe_id(_first(source, "baseline_id", "id"), f"BASE_{index:03d}")
        common = {
            "characteristic": characteristic,
            "subgroup_type": subgroup_type,
            "subgroup_value": subgroup_value,
            "independent_study_id": study_id,
            "comparison_id": None,
            "timepoint": 0,
            "time_unit": "baseline",
            "arm_id": _first(source, "arm_id"),
            "n_baseline": _first(source, "n_baseline", "sample_size", "n"),
            "analysis_population": "baseline",
            "denominator_type": "baseline_as_reported",
            "source_status": status,
            "conflict_id": None,
            "canonical_for_analysis": True,
            "meta_analysis_eligibility": "NOT_POOLABLE",
            "meta_analysis_reason": "Baseline descriptive record; do not pool as an intervention effect.",
            "include_in_meta_analysis": False,
            "source_page": page,
            "source_location": _first(source, "source_location"),
        }
        if _first(source, "mean") is not None:
            result.append({**common, "id": f"{base_id}_MEAN", "canonical_record_id": f"{base_id}_MEAN", "estimand_family": "baseline_mean", "measure": "mean", "estimate_type": "mean", "value": _first(source, "mean"), "unit": _first(source, "unit"), "dispersion": "SD" if _first(source, "sd") is not None else None, "dispersion_value": _first(source, "sd"), "lower_value": None, "upper_value": None, "event_count": None, "denominator": None, "notes": None})
        if _first(source, "median") is not None:
            result.append({**common, "id": f"{base_id}_MEDIAN", "canonical_record_id": f"{base_id}_MEDIAN", "estimand_family": "baseline_median", "measure": "median", "estimate_type": "median", "value": _first(source, "median"), "unit": _first(source, "unit"), "dispersion": "IQR" if _first(source, "q1", "q3") is not None else None, "dispersion_value": None, "lower_value": _first(source, "q1", "min"), "upper_value": _first(source, "q3", "max"), "event_count": None, "denominator": None, "notes": None})
        count = _first(source, "count")
        percentage = _first(source, "percentage")
        baseline_n = common["n_baseline"]
        if percentage is not None:
            result.append({**common, "id": f"{base_id}_PROP", "canonical_record_id": f"{base_id}_PROP", "estimand_family": "baseline_proportion", "measure": "sample_proportion", "estimate_type": "sample_proportion", "value": percentage, "unit": "%", "percentage_unit": "%", "dispersion": None, "dispersion_value": None, "lower_value": None, "upper_value": None, "event_count": count, "event_count_unit": "participants" if count is not None else None, "denominator": baseline_n, "event_denominator": baseline_n, "notes": "Printed count and sample proportion are stored in separate cells on the same observation row." if count is not None else None})
        elif count is not None:
            result.append({**common, "id": f"{base_id}_COUNT", "canonical_record_id": f"{base_id}_COUNT", "estimand_family": "baseline_count", "measure": "count", "estimate_type": "count", "value": count, "unit": "participants", "dispersion": None, "dispersion_value": None, "lower_value": None, "upper_value": None, "event_count": count, "event_count_unit": "participants", "denominator": baseline_n, "event_denominator": baseline_n, "percentage_unit": None, "notes": None})
    return result


def _outcome_measure(source: dict[str, Any]) -> tuple[str, Any, str | None, Any, Any, Any]:
    estimate = str(_first(source, "estimate_type") or "reported_observation").strip().lower()
    value = _first(source, "value", "change_value", "mean", "median", "percentage", "event_count", "events")
    dispersion = "SD" if _first(source, "sd", "change_sd") is not None else ("SE" if _first(source, "se", "change_se") is not None else None)
    dispersion_value = _first(source, "sd", "change_sd", "se", "change_se")
    lower = _first(source, "ci_lower", "iqr_lower")
    upper = _first(source, "ci_upper", "iqr_upper")
    if "percentage" in estimate or _first(source, "percentage") is not None:
        return "sample_proportion", _first(source, "percentage", "value"), None, None, lower, upper
    if "event" in estimate:
        return "event_count", _first(source, "event_count", "events", "value"), None, None, lower, upper
    if "hazard" in estimate:
        return "hazard_ratio", value, None, None, lower, upper
    if "odds" in estimate:
        return "odds_ratio", value, None, None, lower, upper
    if "log" in estimate and "percent" in estimate:
        return "log_transformed_relative_percentage_change", value, None, None, lower, upper
    if "ratio" in estimate:
        return "ratio", value, None, None, lower, upper
    if "difference" in estimate:
        return "mean_difference", value, None, None, lower, upper
    if "mean" in estimate and "change" in estimate:
        return "mean_change", value, dispersion, dispersion_value, lower, upper
    if "mean" in estimate:
        return "mean", value, dispersion, dispersion_value, lower, upper
    if "median" in estimate:
        return "median", value, "IQR" if lower is not None or upper is not None else None, None, lower, upper
    return estimate or "reported_observation", value, dispersion, dispersion_value, lower, upper


def _outcomes_and_tests(
    rows: list[dict[str, Any]], provenance: dict[str, dict[str, Any]], pages: list[str], study_id: str
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    outcomes: list[dict[str, Any]] = []
    tests_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}
    for index, source in enumerate(rows, start=1):
        raw_outcome_name = _first(source, "name", "parameter_name") or f"outcome_{index}"
        outcome_name, threshold_lt, threshold_eq, threshold_gt, threshold_value, threshold_unit = _split_outcome_threshold(raw_outcome_name)
        source_status, page, _ = _verification_for_record(source, provenance, pages, outcome_name)
        outcome_id = _safe_id(_first(source, "outcome_id", "id"), f"OUT_{index:04d}")
        domain = str(_first(source, "domain") or "efficacy").strip().lower()
        measure, value, dispersion, dispersion_value, lower, upper = _outcome_measure(source)
        event_count = _first(source, "event_count", "events")
        denominator = _first(source, "denominator", "total")
        n_analyzed = _first(source, "n_analyzed", "sample_size")
        ci_lower = _first(source, "ci_lower", "ratio_ci_lower")
        ci_upper = _first(source, "ci_upper", "ratio_ci_upper")
        iqr_lower = _first(source, "iqr_lower", "q1")
        iqr_upper = _first(source, "iqr_upper", "q3")
        iqr_duplicated_as_ci = (
            measure == "median"
            and iqr_lower is not None and iqr_upper is not None
            and ci_lower == iqr_lower and ci_upper == iqr_upper
            and not _first(source, "effect_measure", "ci_reported_separately")
        )
        if iqr_duplicated_as_ci:
            ci_lower = None
            ci_upper = None
        effect_measure = _first(source, "effect_measure")
        effect_value = _first(source, "effect_value")
        if measure in {"hazard_ratio", "odds_ratio", "ratio", "mean_difference", "log_transformed_relative_percentage_change"}:
            effect_measure = effect_measure or measure
            effect_value = effect_value if effect_value is not None else value
            ci_lower = ci_lower if ci_lower is not None else lower
            ci_upper = ci_upper if ci_upper is not None else upper
        complete_continuous = measure in {"mean", "mean_change"} and dispersion == "SD" and dispersion_value is not None and n_analyzed is not None
        complete_binary = measure in {"sample_proportion", "event_count"} and event_count is not None and denominator is not None
        complete_effect = effect_measure is not None and effect_value is not None and ci_lower is not None and ci_upper is not None
        approximate = _truth(_first(source, "approximate_from_figure")) or str(_first(source, "derivation_type") or "").lower().startswith("digit")
        reason = "Required arm result, denominator and uncertainty fields are present."
        if approximate:
            eligibility, reason, include = "NOT_POOLABLE", "Figure-derived value was not independently digitized and verified.", False
        else:
            eligibility, reason, include = _analysis_status(source_status, complete_continuous or complete_binary or complete_effect, reason if (complete_continuous or complete_binary or complete_effect) else "The reported result lacks a required denominator, analyzed N, variance, or confidence interval.")
        p_operator, p_value = _parse_p_value(_first(source, "p_value"), _first(source, "p_operator"))
        comparator = _first(source, "comparator_arm_id")
        comparison_seed = "|".join(str(item or "") for item in (outcome_name, _first(source, "timepoint_value", "Timepoint (value)"), _first(source, "timepoint_unit", "Timepoint (unit)"), p_operator, p_value))
        comparison_id = f"CMP_{hashlib.sha256(comparison_seed.encode('utf-8')).hexdigest()[:12].upper()}" if p_value is not None or comparator else None
        record = {
            "id": outcome_id,
            "domain": domain,
            "outcome": outcome_name,
            "independent_study_id": study_id,
            "estimand_family": measure,
            "threshold_lt": threshold_lt,
            "threshold_eq": threshold_eq,
            "threshold_gt": threshold_gt,
            "threshold_value": threshold_value,
            "threshold_unit": threshold_unit,
            "subgroup_type": "reported_subgroup" if _first(source, "definition") else None,
            "subgroup_value": _first(source, "definition"),
            "adherence_level": None,
            "comparison_id": comparison_id,
            "period_id": None,
            "timepoint": _first(source, "timepoint_value", "Timepoint (value)"),
            "time_unit": _first(source, "timepoint_unit", "Timepoint (unit)"),
            "arm_id": _first(source, "arm_id"),
            "n_randomized": _first(source, "n_randomized"),
            "n_baseline": _first(source, "n_baseline"),
            "n_analyzed": n_analyzed,
            "n_analyzed_timepoint": _first(source, "timepoint_value", "Timepoint (value)"),
            "event_count": event_count,
            "event_count_unit": "participants" if event_count is not None else None,
            "denominator": denominator,
            "event_denominator": denominator,
            "percentage_unit": "%" if measure == "sample_proportion" else None,
            "denominator_type": "outcome_specific",
            "measure": measure,
            "estimate_type": _first(source, "estimate_type") or measure,
            "value": value,
            "unit": "%" if measure == "sample_proportion" else ("participants" if measure == "event_count" else _first(source, "unit")),
            "dispersion": dispersion,
            "dispersion_value": dispersion_value,
            "lower_value": lower if measure == "median" else None,
            "upper_value": upper if measure == "median" else None,
            "ci_lower": ci_lower,
            "ci_upper": ci_upper,
            "ci_level": _first(source, "ci_level") or (95 if ci_lower is not None and ci_upper is not None else None),
            "effect_measure": effect_measure,
            "effect_value": effect_value,
            "analysis_population": _first(source, "analysis_population"),
            "applicability_status": "reported",
            "derivation_type": "figure_approximation" if approximate else _first(source, "derivation_type"),
            "source_status": source_status,
            "conflict_id": None,
            "canonical_record_id": outcome_id,
            "canonical_for_analysis": True,
            "meta_analysis_eligibility": eligibility,
            "meta_analysis_reason": reason,
            "include_in_meta_analysis": include,
            "notes": ("IQR values were removed from confidence-interval fields; no separate CI was published." if iqr_duplicated_as_ci else _first(source, "original_timepoint_label")),
            "source_location": _first(source, "source_location"),
            "source_page": page,
        }
        outcomes.append(record)
        if p_value is not None:
            key = (comparison_id, p_operator, p_value)
            test = tests_by_key.setdefault(key, {
                "test_id": f"TEST_{comparison_id}",
                "comparison_id": comparison_id,
                "test_scope": "between_arm_comparison" if comparator else "reported_comparison",
                "domain": domain,
                "outcome": outcome_name,
                "independent_study_id": study_id,
                "estimand_family": "p_value_only",
                "threshold_lt": threshold_lt,
                "threshold_eq": threshold_eq,
                "threshold_gt": threshold_gt,
                "threshold_value": threshold_value,
                "threshold_unit": threshold_unit,
                "period_id": None,
                "timepoint": record["timepoint"],
                "time_unit": record["time_unit"],
                "arm_a_id": record["arm_id"],
                "arm_b_id": comparator,
                "arm_id": None,
                "arm_a_outcome_id": outcome_id,
                "arm_b_outcome_id": None,
                "subgroup_factor": record["subgroup_value"],
                "test_effect": None,
                "estimate_type": "p_value",
                "effect_measure": effect_measure,
                "effect_value": effect_value,
                "ci_lower": ci_lower,
                "ci_upper": ci_upper,
                "ci_level": record["ci_level"],
                "analysis_population": record["analysis_population"],
                "p_lt": p_operator == "<",
                "p_eq": p_operator == "=",
                "p_gt": p_operator == ">",
                "p_value": p_value,
                "nominal_adjusted_status": _first(source, "p_value_status", "nominal_adjusted_status"),
                "source_status": source_status,
                "conflict_id": None,
                "canonical_for_analysis": True,
                "meta_analysis_eligibility": "READY" if include and complete_effect else "NOT_POOLABLE",
                "meta_analysis_reason": "Published effect estimate and confidence interval are available." if include and complete_effect else "P-value-only comparison; retain for interpretation but do not pool as an effect estimate.",
                "include_in_meta_analysis": bool(include and complete_effect),
                "source_location": _first(source, "source_location"),
                "source_page": page,
                "notes": "Stored once for the comparison; not duplicated on arm-level outcome rows.",
            })
            if test["arm_a_id"] != record["arm_id"] and not test.get("arm_b_id"):
                test["arm_b_id"] = record["arm_id"]
                test["arm_b_outcome_id"] = outcome_id
    return outcomes, list(tests_by_key.values())


def _arm_alias_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())


def _canonicalize_arm_references(
    arms: list[dict[str, Any]],
    baseline: list[dict[str, Any]],
    outcomes: list[dict[str, Any]],
    statistical_tests: list[dict[str, Any]],
) -> tuple[int, int]:
    """Resolve common Mandros arm aliases without guessing across equal arms.

    Some exports use canonical IDs in ``Arms`` but short semantic IDs such as
    ``ARM_PLA`` or ``ARM_DAPA_10`` in outcome sheets. We resolve an alias only
    when its label/drug/dose signature or sample size selects exactly one arm.
    """

    aliases: dict[str, set[str]] = {}
    size_index: dict[int, set[str]] = {}
    for arm in arms:
        arm_id = arm["arm_id"]
        candidates = {arm_id, arm.get("label"), arm.get("drug")}
        drug = str(arm.get("drug") or "").strip()
        label = str(arm.get("label") or "").strip()
        dose = arm.get("dose_value")
        if "placebo" in f"{drug} {label}".lower():
            candidates.update({"ARM_PLA", "ARM_PBO", "PLA", "PBO", "PLACEBO"})
        if drug and dose is not None:
            prefix = re.sub(r"[^A-Za-z0-9]", "", drug)[:4]
            candidates.update({f"ARM_{prefix}_{dose}", f"{prefix}_{dose}", f"{drug}_{dose}", f"{label}_{dose}"})
        for candidate in candidates:
            key = _arm_alias_key(candidate)
            if key:
                aliases.setdefault(key, set()).add(arm_id)
        for size in (arm.get("n_randomized"), arm.get("n_baseline"), arm.get("n_analyzed")):
            if isinstance(size, (int, float)) and int(size) == size:
                size_index.setdefault(int(size), set()).add(arm_id)

    remapped = 0
    unresolved = 0

    def resolve(value: Any, sample_size: Any = None) -> Any:
        nonlocal remapped, unresolved
        if not value or any(value == arm["arm_id"] for arm in arms):
            return value
        candidates = set(aliases.get(_arm_alias_key(value), set()))
        if isinstance(sample_size, (int, float)) and int(sample_size) == sample_size:
            size_candidates = size_index.get(int(sample_size), set())
            candidates = candidates & size_candidates if candidates and size_candidates else (candidates or set(size_candidates))
        if len(candidates) == 1:
            remapped += 1
            return next(iter(candidates))
        unresolved += 1
        return value

    for record in baseline:
        original = record.get("arm_id")
        record["arm_id"] = resolve(original, record.get("n_baseline"))
        if original and record["arm_id"] != original:
            record["notes"] = f"Arm alias {original} normalized to {record['arm_id']}."
    for record in outcomes:
        original = record.get("arm_id")
        record["arm_id"] = resolve(original, record.get("n_analyzed") or record.get("denominator"))
        if original and record["arm_id"] != original:
            note = f"Arm alias {original} normalized to {record['arm_id']}."
            record["notes"] = f"{record['notes']} {note}".strip() if record.get("notes") else note
    for test in statistical_tests:
        test["arm_a_id"] = resolve(test.get("arm_a_id"))
        test["arm_b_id"] = resolve(test.get("arm_b_id"))
        test["arm_id"] = resolve(test.get("arm_id"))
    return remapped, unresolved


def build_generic_dataset(pdf_path: Path, workbook_path: Path, validation_report: dict[str, Any]) -> dict[str, Any]:
    pages = extract_pages(pdf_path)
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        provenance_rows = _records(workbook, "Source Provenance") or _records(workbook, "Provenance")
        provenance = {
            str(_first(row, "source_id", "provenance_id")): row
            for row in provenance_rows
            if _first(row, "source_id", "provenance_id")
        }
        study = _study_record(_records(workbook, "Studies"), pdf_path, pages)
        arms_source = _records(workbook, "Treatment Arms") or _records(workbook, "Arms")
        eligibility_source = _records(workbook, "Eligibility Criteria") or _records(workbook, "Eligibility")
        baseline_source = _records(workbook, "Baseline Characteristics") or _records(workbook, "Baseline")
        outcome_source = _records(workbook, "Efficacy Outcomes") + _records(workbook, "Safety Outcomes")
        if not outcome_source:
            outcome_source = [row for row in _records(workbook, "Outcomes") if not _truth(row.get("is_baseline"))]
        arms = _arms(arms_source, provenance, pages)
        eligibility = _eligibility(eligibility_source, provenance, pages)
        baseline = _baseline(baseline_source, provenance, pages, study["study_id"])
        outcomes, statistical_tests = _outcomes_and_tests(outcome_source, provenance, pages, study["study_id"])
    finally:
        workbook.close()

    remapped_arm_records, _ = _canonicalize_arm_references(arms, baseline, outcomes, statistical_tests)
    valid_arm_ids = {item["arm_id"] for item in arms}
    invalid_arm_records = 0
    for item in [*baseline, *outcomes]:
        if item.get("arm_id") and item["arm_id"] not in valid_arm_ids:
            invalid_arm_records += 1
            item["meta_analysis_eligibility"] = "REQUIRES_REVIEW"
            item["meta_analysis_reason"] = "The arm_id does not resolve to a canonical arm in the supplied Mandros workbook."
            item["include_in_meta_analysis"] = False
            item["notes"] = "Invalid arm reference; resolve against the PDF before analysis."

    all_records = [*arms, *eligibility, *baseline, *outcomes]
    exact = sum(item.get("source_status") == "published_main_article_exact" for item in all_records)
    review = sum(item.get("source_status") != "published_main_article_exact" for item in all_records)
    issue_rows = [
        {
            "code": issue.get("code"),
            "status": "OPEN" if issue.get("blocking") else "MANUAL_REVIEW",
            "severity": issue.get("severity"),
            "message": issue.get("message"),
        }
        for issue in validation_report.get("issues", [])
    ]
    issue_rows.extend([
        {"code": "GENERIC_PDF_MANDROS_PIPELINE", "status": "RESOLVED", "severity": "information", "message": "Workbook generated from the supplied PDF and Mandros Excel only; no manual workbook was used."},
        {"code": "SOURCE_GROUNDING", "status": "MANUAL_REVIEW" if review else "RESOLVED", "severity": "major" if review else "information", "message": f"Exact page-level evidence matches: {exact}; records requiring page-level review: {review}."},
        {"code": "ARM_REFERENCE_CHECK", "status": "MANUAL_REVIEW" if invalid_arm_records else "RESOLVED", "severity": "critical" if invalid_arm_records else "information", "message": f"Outcome/baseline records with an unresolved arm_id: {invalid_arm_records}."},
        {"code": "ARM_ALIAS_NORMALIZATION", "status": "RESOLVED", "severity": "information", "message": f"Arm aliases normalized to canonical Arms IDs: {remapped_arm_records}."},
        {"code": "COMPLETENESS_LIMIT", "status": "MANUAL_REVIEW", "severity": "major", "message": "The generic pass normalizes and verifies Mandros records but cannot prove that Mandros extracted every result printed in an unseen PDF."},
    ])
    output_name = f"{_safe_id(workbook_path.stem, 'study').lower()}_Mandros_Blind_corrected.xlsx"
    return {
        "dataset_id": f"generic-{sha256_file(pdf_path)[:12]}",
        "version": "1.0.0",
        "generation_mode": "generic_pdf_mandros",
        "output_filename": output_name,
        "preview_prefix": f"generic-{sha256_file(pdf_path)[:10]}",
        "source_document": pdf_path.name,
        "source_sha256": sha256_file(pdf_path),
        "mandros_document": workbook_path.name,
        "mandros_sha256": sha256_file(workbook_path),
        "document_inventory": [
            {"document_id": "DOC_PDF", "document_type": "main_article_pdf", "document_name": pdf_path.name, "availability_status": "supplied", "extraction_status": "read", "page_count": len(pages), "sha256": sha256_file(pdf_path)},
            {"document_id": "DOC_MANDROS", "document_type": "mandros_extraction_xlsx", "document_name": workbook_path.name, "availability_status": "supplied", "extraction_status": "normalized", "page_count": None, "sha256": sha256_file(workbook_path)},
        ],
        "study": study,
        "master_study_rows": [
            ["Study ID", study["study_id"], "DOI", study.get("doi"), "Year", study.get("year"), "Randomized total N", study.get("randomized_total_n")],
            ["Design", study.get("design"), "Study phase", study.get("study_phase"), "Duration", study.get("duration_value"), "Unit", study.get("duration_unit")],
            ["Title", study.get("title"), "Condition", study.get("condition"), "Source", pdf_path.name, "Input", workbook_path.name],
        ],
        "arms": arms,
        "eligibility": eligibility,
        "baseline": baseline,
        "baseline_status": "Baseline records normalized from the Mandros workbook." if baseline else "No baseline records were present in the Mandros workbook; none were invented.",
        "outcomes": outcomes,
        "statistical_tests": statistical_tests,
        "source_conflicts": [],
        "phase2_readiness": [
            {"row_type": "input", "item": "Main article PDF", "phase_1_requirement": "required", "extraction_status": "supplied", "phase_2_use": "source verification", "meta_analysis_eligibility": "REQUIRES_REVIEW", "source_page": None, "notes": pdf_path.name},
            {"row_type": "input", "item": "Mandros Excel", "phase_1_requirement": "required", "extraction_status": "normalized", "phase_2_use": "candidate records", "meta_analysis_eligibility": "REQUIRES_REVIEW", "source_page": None, "notes": workbook_path.name},
            {"row_type": "gate", "item": "Full-PDF completeness", "phase_1_requirement": "independent confirmation", "extraction_status": "not_proven_by_generic_pass", "phase_2_use": "prevents silent outcome omission", "meta_analysis_eligibility": "REQUIRES_REVIEW", "source_page": None, "notes": "No manual workbook is required, but an unseen-study generic pass cannot prove exhaustive extraction without a trained semantic extraction model or human adjudication."},
        ],
        "statistical_methods": [],
        "validation_log": issue_rows,
        "qa_rows": [
            ["Generation mode", None, None, "PDF + Mandros Excel only"],
            ["Manual workbook used", 0, 0, "NO"],
            ["Arm records", len(arms), None, "NORMALIZED"],
            ["Eligibility records", len(eligibility), None, "NORMALIZED"],
            ["Baseline records", len(baseline), None, "NORMALIZED; ATOMIC"],
            ["Outcome records", len(outcomes), None, "P-VALUES REMOVED FROM ARM ROWS"],
            ["Statistical test records", len(statistical_tests), None, "STORED ONCE PER COMPARISON"],
            ["Exact source matches", exact, None, "SOURCE-CHECKED"],
            ["Records requiring source review", review, None, "REQUIRES_REVIEW"],
            ["Scientific status", None, None, "STRUCTURALLY CORRECTED; REVIEW FLAGS PRESERVED"],
        ],
    }
