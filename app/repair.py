from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .metformin_data import load_metformin_bundle


NAVY = "17324D"
TEAL = "0F766E"
PALE = "E8F3F1"
WHITE = "FFFFFF"


def _style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        letter = column[0].column_letter
        width = min(32, max(11, max(len(str(c.value or "")) for c in column) + 2))
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _append_dicts(ws, rows: list[dict[str, Any]], headers: list[str]) -> None:
    ws.append(headers)
    for record in rows:
        ws.append([record.get(h) for h in headers])
    _style_sheet(ws)


def _operator_symbol(value: Any, symbol: str) -> str | None:
    return symbol if value is True else None


def _boundary_label(value: Any) -> str | None:
    if value is True:
        return "inclusive"
    if value is False:
        return "exclusive"
    return None


def _append_projected_dicts(
    ws,
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str, Any]],
) -> None:
    ws.append([label for _, label, _ in columns])
    for record in rows:
        values = []
        for key, _, transform in columns:
            value = record.get(key)
            values.append(transform(value) if transform else value)
        ws.append(values)
    _style_sheet(ws)


def create_corrected_workbook(output_path: Path) -> None:
    bundle = load_metformin_bundle()
    wb = Workbook()
    wb.remove(wb.active)

    master = wb.create_sheet("Master Extraction")
    master.sheet_view.showGridLines = False
    master.merge_cells("A1:H1")
    master["A1"] = "MANDROS BLIND — SOURCE-GROUNDED MASTER EXTRACTION"
    master["A1"].fill = PatternFill("solid", fgColor=NAVY)
    master["A1"].font = Font(color=WHITE, bold=True, size=15)
    master["A1"].alignment = Alignment(horizontal="center")
    rows = [
        ("Validation status", "Corrected against supplied PDF; baseline remains unreported rather than fabricated"),
        ("Study ID", bundle["study"]["study_id"]),
        ("Title", bundle["study"]["title"]),
        ("DOI", bundle["study"]["doi"]),
        ("Randomized total N", bundle["study"]["randomized_total_n"]),
        ("Intervention", "Metformin 850 mg orally once daily; target twice daily after 1 month as tolerated"),
        ("Comparator", "Matching placebo on the same schedule during the randomized blinded phase"),
        ("Baseline", bundle["baseline_status"]),
        ("Efficacy records", sum(1 for x in bundle["outcomes"] if x["domain"] == "efficacy")),
        ("Safety records", sum(1 for x in bundle["outcomes"] if x["domain"] == "safety")),
        ("Statistical test records", len(bundle["statistical_tests"])),
        ("Eligibility records", len(bundle["eligibility"])),
        ("Rule", "One row = one observation; value/unit/operator/timepoint are separate fields")
    ]
    master.append(["Field", "Value"])
    for r in rows:
        master.append(list(r))
    master["A2"].fill = PatternFill("solid", fgColor=TEAL)
    master["B2"].fill = PatternFill("solid", fgColor=TEAL)
    for c in master[2]:
        c.font = Font(color=WHITE, bold=True)
    master.column_dimensions["A"].width = 27
    master.column_dimensions["B"].width = 95
    for row in master.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    studies = wb.create_sheet("Studies")
    _append_dicts(studies, [bundle["study"]], list(bundle["study"].keys()))

    arms = wb.create_sheet("Arms")
    _append_dicts(arms, bundle["arms"], list(bundle["arms"][0].keys()))

    elig = wb.create_sheet("Eligibility")
    eligibility_columns = [
        ("criterion_id", "criterion_id", None),
        ("category", "category", None),
        ("variable", "variable", None),
        ("operator_lt", "<", lambda value: _operator_symbol(value, "<")),
        ("operator_eq", "=", lambda value: _operator_symbol(value, "=")),
        ("operator_gt", ">", lambda value: _operator_symbol(value, ">")),
        ("value", "value", None),
        ("unit", "unit", None),
        ("range_min", "range_min", None),
        ("range_max", "range_max", None),
        ("range_unit", "range_unit", None),
        ("range_min_inclusive", "range_min_boundary", _boundary_label),
        ("range_max_inclusive", "range_max_boundary", _boundary_label),
        ("population_scope", "population_scope", None),
        ("qualifier", "qualifier", None),
        ("source_page", "source_page", None),
    ]
    _append_projected_dicts(elig, bundle["eligibility"], eligibility_columns)

    baseline = wb.create_sheet("Baseline")
    _append_dicts(baseline, [{"study_id": bundle["study"]["study_id"], "status": "not_reported_in_this_article", "note": bundle["baseline_status"], "source_page": None}], ["study_id", "status", "note", "source_page"])

    outcomes = wb.create_sheet("Outcomes")
    outcome_columns = [
        ("id", "id", None), ("domain", "domain", None), ("outcome", "outcome", None),
        ("threshold_lt", "threshold <", lambda value: _operator_symbol(value, "<")),
        ("threshold_eq", "threshold =", lambda value: _operator_symbol(value, "=")),
        ("threshold_gt", "threshold >", lambda value: _operator_symbol(value, ">")),
        ("threshold_value", "threshold_value", None), ("threshold_unit", "threshold_unit", None),
        ("subgroup_type", "subgroup_type", None), ("subgroup_value", "subgroup_value", None),
        ("adherence_level", "adherence_level", None),
        ("comparison_id", "comparison_id", None),
        ("period_id", "period_id", None), ("timepoint", "timepoint", None),
        ("time_unit", "time_unit", None), ("arm_id", "arm_id", None),
        ("measure", "measure", None), ("value", "value", None), ("unit", "unit", None),
        ("dispersion", "dispersion", None), ("dispersion_value", "dispersion_value", None),
        ("adherence_min_value", "adherence_min_value", None),
        ("adherence_min_unit", "adherence_min_unit", None),
        ("adherence_min_inclusive", "adherence_min_boundary", _boundary_label),
        ("adherence_max_value", "adherence_max_value", None),
        ("adherence_max_unit", "adherence_max_unit", None),
        ("adherence_max_inclusive", "adherence_max_boundary", _boundary_label),
        ("visit_pill_lt", "visit_pill <", lambda value: _operator_symbol(value, "<")),
        ("visit_pill_eq", "visit_pill =", lambda value: _operator_symbol(value, "=")),
        ("visit_pill_gt", "visit_pill >", lambda value: _operator_symbol(value, ">")),
        ("visit_pill_min_value", "visit_pill_min_value", None),
        ("visit_pill_min_unit", "visit_pill_min_unit", None),
        ("adherence_special_rule", "adherence_special_rule", None),
        ("analysis_population", "analysis_population", None),
        ("applicability_status", "applicability_status", None),
        ("source_page", "source_page", None),
    ]
    _append_projected_dicts(outcomes, bundle["outcomes"], outcome_columns)

    tests = wb.create_sheet("Statistical Tests")
    statistical_test_columns = [
        ("test_id", "test_id", None), ("comparison_id", "comparison_id", None),
        ("test_scope", "test_scope", None), ("domain", "domain", None), ("outcome", "outcome", None),
        ("threshold_lt", "threshold <", lambda value: _operator_symbol(value, "<")),
        ("threshold_eq", "threshold =", lambda value: _operator_symbol(value, "=")),
        ("threshold_gt", "threshold >", lambda value: _operator_symbol(value, ">")),
        ("threshold_value", "threshold_value", None), ("threshold_unit", "threshold_unit", None),
        ("period_id", "period_id", None), ("timepoint", "timepoint", None),
        ("time_unit", "time_unit", None), ("arm_a_id", "arm_a_id", None),
        ("arm_b_id", "arm_b_id", None), ("arm_id", "arm_id", None),
        ("arm_a_outcome_id", "arm_a_outcome_id", None),
        ("arm_b_outcome_id", "arm_b_outcome_id", None),
        ("subgroup_factor", "subgroup_factor", None),
        ("p_lt", "p <", lambda value: _operator_symbol(value, "<")),
        ("p_eq", "p =", lambda value: _operator_symbol(value, "=")),
        ("p_gt", "p >", lambda value: _operator_symbol(value, ">")),
        ("p_value", "p_value", None), ("source_page", "source_page", None),
        ("notes", "notes", None),
    ]
    _append_projected_dicts(tests, bundle["statistical_tests"], statistical_test_columns)

    provenance = wb.create_sheet("Provenance")
    prov_rows = [
        *[{"record_id": x["id"], "record_type": "outcome", "source_document": "metformin clinical trial.pdf", "source_page": x["source_page"], "verification_status": "source_checked", "evidence_type": "published_text_or_table"} for x in bundle["outcomes"]],
        *[{"record_id": x["test_id"], "record_type": "statistical_test", "source_document": "metformin clinical trial.pdf", "source_page": x["source_page"], "verification_status": "source_checked", "evidence_type": "published_text_or_table"} for x in bundle["statistical_tests"]],
    ]
    _append_dicts(provenance, prov_rows, ["record_id", "record_type", "source_document", "source_page", "verification_status", "evidence_type"])

    validation = wb.create_sheet("Validation Log")
    log_rows = [
        {"code": "FIX_ARM_IDS", "status": "RESOLVED", "severity": "critical", "message": "Canonical ARM_MET and ARM_PBO are used throughout."},
        {"code": "FIX_SAMPLE_SCOPE", "status": "RESOLVED", "severity": "critical", "message": "Year-9 visit counts 790/783 are not used as randomized or universal analyzed N."},
        {"code": "FIX_PLACEBO_Y2", "status": "RESOLVED", "severity": "critical", "message": "Year-2 placebo weight loss corrected to 0.02 +/- 5.5%."},
        {"code": "REMOVE_FABRICATED_BASELINE", "status": "RESOLVED", "severity": "critical", "message": "Fabricated 100% Caucasian baseline rows removed."},
        {"code": "ADD_SAFETY", "status": "RESOLVED", "severity": "critical", "message": "Published safety results added as arm-specific observations."},
        {"code": "ADD_ALL_ADHERENCE_LEVELS", "status": "RESOLVED", "severity": "major", "message": "All four Table 1 adherence levels were added across applicable arms, periods and subgroups."},
        {"code": "SOURCE_UNIT_ANOMALY", "status": "MANUAL_REVIEW", "severity": "major", "message": "The article prints hemoglobin in mg/dL; preserve as published but verify against protocol or erratum before pooling."},
        {"code": "BASELINE_SCOPE", "status": "OPEN", "severity": "information", "message": bundle["baseline_status"]},
        {"code": "FULL_PDF_SCOPE", "status": "MANUAL_REVIEW", "severity": "major", "message": "This corrected workbook covers the primary efficacy, dosing, eligibility and key safety results. Independent full-PDF adjudication remains required before publication or meta-analysis."}
    ]
    _append_dicts(validation, log_rows, ["code", "status", "severity", "message"])

    qa = wb.create_sheet("Quality Assessment")
    _append_dicts(qa, [
        {"metric": "Original Mandros score", "value": 38, "maximum": 100, "decision": "REJECT"},
        {"metric": "Manual benchmark score", "value": 55, "maximum": 100, "decision": "NOT COMPLETE"},
        {"metric": "Corrected workbook state", "value": None, "maximum": None, "decision": "SOURCE-CHECKED CORE; MANUAL FULL-SCOPE REVIEW REQUIRED"}
    ], ["metric", "value", "maximum", "decision"])

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = TEAL
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

